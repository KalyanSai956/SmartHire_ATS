import io
import os
import re

from typing import Tuple, Optional, List

from backend.models.schemas import (
    ResumeProfile,
    ResumeContact,
    ResumeEducation,
    ResumeExperience,
    ResumeProject,
    ResumeCertification,
    ResumeAchievement,
)

import pdfplumber
from docx import Document
import PyPDF2
from backend.utils.file_utils import (
    FileParsingError,
    TextExtractionError,
    FileUploadError,
    log_error,
    log_warning,log_info,
    with_fallback
)
from backend.core.config import (
    MAX_FILE_SIZE_MB,
    SUPPORTED_MIME_TYPES,
    MAX_FILE_SIZE_BYTES
)
class FileValidationError(Exception):
    pass

def validate_file(file_data: bytes, filename: str) -> Tuple[bool, str, Optional[str]]:
    file_size_bytes = len(file_data)

    if file_size_bytes > MAX_FILE_SIZE_BYTES:
        size_mb = file_size_bytes / (1024 * 1024)
        return (
            False,
            f"File size ({size_mb:.2f} MB) exceeds the maximum of {MAX_FILE_SIZE_MB} MB.",
            None,
        )

    if file_size_bytes == 0:
        return False, "Uploaded file is empty.", None

    extension = os.path.splitext(filename)[1].lower()

    extension_map = {
        ".pdf": "pdf",
        ".doc": "doc",
        ".docx": "docx",
    }

    if extension not in extension_map:
        return (
            False,
            "Unsupported file type. Please upload a PDF, DOC, or DOCX file.",
            None,
        )

    return True, "", extension_map[extension]
def _extract_pdf_hyperlinks(file_data: bytes) -> str:
    urls = []
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_data))
        for page in reader.pages:
            if '/Annots' not in page:
                continue
            for annot_ref in page['/Annots']:
                try:
                    annot = annot_ref.get_object()
                    if annot.get('/Subtype') != '/Link':
                        continue
                    action = annot.get('/A', {})
                    uri = action.get('/URI', '')
                    if uri and isinstance(uri, (str, bytes)):
                        # PyPDF2 may return bytes for URI values
                        if isinstance(uri, bytes):
                            uri = uri.decode('utf-8', errors='ignore')
                        uri = uri.strip()
                        if uri.startswith('http'):
                            urls.append(uri)
                except Exception:
                    pass
    except Exception:
        pass
    return '\n'.join(urls)


def _extract_pdf_with_pdfplumber(file_data: bytes) -> str:
    text = ''
    with pdfplumber.open(io.BytesIO(file_data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + '\n'

    if not text.strip():
        raise TextExtractionError(
            'pdfplumber extracted no text',
            user_message='No text could be extracted from the PDF.'
        )
    
    hyperlinks = _extract_pdf_hyperlinks(file_data)
    if hyperlinks:
        text = text.strip() + '\n' + hyperlinks

    return text.strip()


def _extract_pdf_with_pypdf2(file_data: bytes) -> str:
    text = ''
    pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_data))
    for page in pdf_reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + '\n'

    if not text.strip():
        raise TextExtractionError(
            'PyPDF2 extracted no text',
            user_message='No text could be extracted from the PDF.'
        )

    hyperlinks = _extract_pdf_hyperlinks(file_data)
    if hyperlinks:
        text = text.strip() + '\n' + hyperlinks

    return text.strip()


def extract_text_from_pdf(file_data: bytes) -> str:
    try: 
        result, used_fallback=with_fallback(
        _extract_pdf_with_pdfplumber, 
        _extract_pdf_with_pypdf2, 
        file_data, 
        log_fallback=True
    )
    
        if used_fallback:
            log_info('PDF EXTRACTION succeded using the PyPDF2 fallback', context='resume_parser')
        return result
        
    except Exception as e:
        log_error(e, context='extract_text_from_pdf')
        raise FileParsingError(
            'Failed to extract text from PDF using both pdfplumber and PyPDF2. '
            'The PDF may be corrupted, password-protected, or contain only scanned images. '
            'Please ensure it contains selectable text.'
        ) from e
    

def extract_text_from_docx(file_data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_data))
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        text = '\n'.join(text_parts)

        if not text.strip():
            raise FileParsingError(
                'No text could be extracted from the document. '
                'The document may be empty or corrupted.'
            )
        
        try:
            for rel in doc.part.rels.values():
                if 'hyperlink' in rel.reltype.lower():
                    url = rel._target
                    if isinstance(url, str) and url.startswith('http'):
                        text += '\n' + url
        except Exception:
            pass

        log_info(f'Extracted {len(text)} chars from DOCX', context='resume_parser')
        return text.strip()

    except FileParsingError:
        raise   # Re-raise unchanged — don't wrap in another FileParsingError

    except Exception as e:
        log_error(e, context='extract_text_from_docx')
        raise FileParsingError(
            'Failed to extract text from DOCX. '
            'The document may be corrupted or in an unsupported format. '
            'Please try re-saving or converting to PDF.'
        ) from e

def extract_text_from_doc(file_data: bytes) -> str:
    raise FileParsingError(
        'Legacy .doc format is not supported. '
        'Please convert your document to .docx or .pdf and try again. '
        'You can convert using Microsoft Word, Google Docs, or online tools.'
    )

def extract_text(file_data:bytes, file_type:str)->str:
    if file_type=='pdf':
        return extract_text_from_pdf(file_data)
    elif file_type=='docx':
        return extract_text_from_docx(file_data)
    elif file_type=='doc':
        return extract_text_from_doc(file_data)
    else:
        raise FileValidationError(
            f'invalid file type: {file_type}. supported types are: pdf, docx and doc'


        )
    
def parse_resume_file(file_data: bytes, filename:str)->Tuple[str, dict]:
    log_info(f'parsing file :{filename}', context='parse_Resume_file')

    #phase01:validate file
    try:
        is_valid, error_msg, file_type=validate_file(file_data, filename)
        if not is_valid:
            log_warning(f'validation failed for file {filename}', context='parse_resume_file')
            raise FileValidationError(error_msg)
    
    except FileValidationError as e:
        raise 

    except Exception as e:
        log_error(e, context='parse_resume_file_validation')
        raise FileValidationError(
            'Could not validate the uploaded file. Please ensure it is a valid PDF or DOCX.'
        ) from e
    
    #phase02: extraction of file

    try:
        text = extract_text(file_data, file_type)
        log_info(f'Extracted {len(text)} chars from {filename}', context='parse_resume_file')

    except FileParsingError:
        raise   # Re-raise unchanged

    except Exception as e:
        log_error(e, context='parse_resume_file_extraction')
        raise FileParsingError(
            'An unexpected error occurred while processing the file. '
            'Please try again or contact support if the problem persists.'
        ) from e

    metadata = {
        'filename':        filename,
        'file_type':       file_type,
        'file_size_bytes': len(file_data),
        'text_length':     len(text),
        'success':         True,
    }
    return text, metadata

# ============================================================
# PHASE 3B — STRUCTURED RESUME PARSING
# ============================================================

SECTION_ALIASES = {
    "summary": {
        "summary",
        "profile",
        "profile summary",
        "professional summary",
        "objective",
        "career objective",
        "about me",
    },
    "education": {
        "education",
        "academic background",
        "academic qualifications",
        "qualifications",
    },
    "experience": {
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "work history",
    },
    "projects": {
        "projects",
        "personal projects",
        "academic projects",
        "project experience",
    },
    "skills": {
        "skills",
        "technical skills",
        "core skills",
        "technologies",
        "technical expertise",
    },
    "certifications": {
        "certifications",
        "certificates",
        "licenses & certifications",
        "licenses and certifications",
    },
    "achievements": {
        "achievements",
        "accomplishments",
        "awards",
        "honors",
        "honours",
    },
}


def _clean_line(line: str) -> str:
    """Normalize whitespace without destroying useful punctuation."""
    return re.sub(r"\s+", " ", line).strip()


def _normalize_heading(line: str) -> str:
    """
    Convert a possible section heading into a normalized form.

    Examples:
        'TECHNICAL SKILLS' -> 'technical skills'
        'Education:'       -> 'education'
    """
    line = line.strip().lower()
    line = re.sub(r"[:|]+$", "", line)
    line = re.sub(r"\s+", " ", line)
    return line


def _detect_section(line: str) -> Optional[str]:
    """
    Detect whether a line represents a known resume section.
    """
    normalized = _normalize_heading(line)

    for section_name, aliases in SECTION_ALIASES.items():
        if normalized in aliases:
            return section_name

    return None


def _split_resume_sections(text: str) -> dict:
    """
    Split raw resume text into logical sections.

    Unknown content before the first recognized heading is kept
    under 'header'.
    """
    sections = {
        "header": [],
        "summary": [],
        "education": [],
        "experience": [],
        "projects": [],
        "skills": [],
        "certifications": [],
        "achievements": [],
    }

    current_section = "header"

    for raw_line in text.splitlines():
        line = _clean_line(raw_line)

        if not line:
            continue

        detected_section = _detect_section(line)

        if detected_section:
            current_section = detected_section
            continue

        sections[current_section].append(line)

    return {
        key: "\n".join(value).strip()
        for key, value in sections.items()
    }


def _extract_email(text: str) -> Optional[str]:
    match = re.search(
        r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text,
    )

    return match.group(0) if match else None


def _extract_phone(text: str) -> Optional[str]:
    match = re.search(
        r"(?<!\d)(?:\+?\d[\d\s().-]{8,}\d)(?!\d)",
        text,
    )

    if not match:
        return None

    phone = re.sub(r"\s+", " ", match.group(0)).strip()

    return phone


def _extract_url(text: str, platform: str) -> Optional[str]:
    pattern = rf"https?://(?:www\.)?{platform}\.[^\s<>,]+"

    match = re.search(pattern, text, re.IGNORECASE)

    if match:
        return match.group(0).rstrip(".,;)")

    # Also support resumes that omit https://
    pattern = rf"(?:www\.)?{platform}\.[^\s<>,]+"

    match = re.search(pattern, text, re.IGNORECASE)

    if match:
        return match.group(0).rstrip(".,;)")

    return None


def _extract_name(text: str) -> Optional[str]:
    """
    Best-effort name extraction.

    Resume parsers should treat this as a heuristic rather than
    assuming the first line is always the candidate's name.
    """
    for line in text.splitlines():
        line = _clean_line(line)

        if not line:
            continue

        if _detect_section(line):
            break

        if "@" in line:
            continue

        if re.search(r"https?://|www\.|linkedin|github", line, re.I):
            continue

        if re.search(r"\d", line):
            continue

        words = line.split()

        if 2 <= len(words) <= 5:
            return line

    return None


def _extract_contact(text: str) -> ResumeContact:
    return ResumeContact(
        name=_extract_name(text),
        email=_extract_email(text),
        phone=_extract_phone(text),
        linkedin=_extract_url(text, "linkedin"),
        github=_extract_url(text, "github"),
        portfolio=None,
    )


def _split_bullets(text: str) -> List[str]:
    """
    Convert bullet-style resume text into clean bullet strings.
    """
    bullets = []

    for line in text.splitlines():
        line = line.strip()

        if not line:
            continue

        line = re.sub(r"^[•●▪◦‣*-]\s*", "", line)
        line = re.sub(r"^\d+[.)]\s*", "", line)

        if line:
            bullets.append(line)

    return bullets


def _parse_skills(section_text: str) -> List[str]:
    """
    Parse common resume skill formats.

    Example:
        Programming Languages: Java, Python, JavaScript
        Frameworks: React.js, FastAPI

    becomes:
        [
            'Java',
            'Python',
            'JavaScript',
            'React.js',
            'FastAPI'
        ]
    """
    skills = []

    for line in section_text.splitlines():
        line = _clean_line(line)

        if not line:
            continue

        if ":" in line:
            _, values = line.split(":", 1)
            candidates = re.split(r",|;|\||•", values)
        else:
            candidates = re.split(r",|;|\||•", line)

        for candidate in candidates:
            candidate = candidate.strip()

            if candidate:
                skills.append(candidate)

    # Preserve order while removing duplicates
    seen = set()
    normalized_skills = []

    for skill in skills:
        key = skill.lower()

        if key not in seen:
            seen.add(key)
            normalized_skills.append(skill)

    return normalized_skills


def _parse_education(section_text: str) -> List[ResumeEducation]:
    if not section_text.strip():
        return []

    lines = [
        _clean_line(line)
        for line in section_text.splitlines()
        if _clean_line(line)
    ]

    if not lines:
        return []

    education = []

    degree_keywords = (
        "bachelor",
        "master",
        "b.tech",
        "m.tech",
        "b.e",
        "m.e",
        "b.sc",
        "m.sc",
        "bca",
        "mca",
        "degree",
        "diploma",
        "ph.d",
        "phd",
    )

    date_pattern = re.compile(
        r"([A-Za-z]+\s+\d{4})\s*[–-]\s*"
        r"([A-Za-z]+\s+\d{4}|Present|Current)",
        re.IGNORECASE,
    )

    grade_pattern = re.compile(
        r"(?:CGPA|GPA|Grade|Percentage)"
        r"\s*[:\-]?\s*([0-9.]+(?:/\d+)?)%?",
        re.IGNORECASE,
    )

    institution = None
    degree = None
    field_of_study = None
    location = None
    start_date = None
    end_date = None
    grade = None

    def flush_education():
        nonlocal institution
        nonlocal degree
        nonlocal field_of_study
        nonlocal location
        nonlocal start_date
        nonlocal end_date
        nonlocal grade

        if not institution and not degree:
            return

        education.append(
            ResumeEducation(
                institution=institution,
                degree=degree,
                field_of_study=field_of_study,
                location=location,
                start_date=start_date,
                end_date=end_date,
                grade=grade,
            )
        )

        institution = None
        degree = None
        field_of_study = None
        location = None
        start_date = None
        end_date = None
        grade = None

    for line in lines:
        lower = line.lower()

        # Date
        date_match = date_pattern.search(line)

        if date_match:
            start_date = date_match.group(1)
            end_date = date_match.group(2)
            continue

        # Grade / CGPA / GPA
        grade_match = grade_pattern.search(line)

        if grade_match:
            grade = grade_match.group(1)
            continue

        # Degree
        if any(keyword in lower for keyword in degree_keywords):
            if degree is None:
                degree_line = line

                # Handle:
                # Bachelor of Technology in Computer Science and Engineering
                if re.search(r"\s+in\s+", degree_line, re.IGNORECASE):
                    degree_part, field_part = re.split(
                        r"\s+in\s+",
                        degree_line,
                        maxsplit=1,
                        flags=re.IGNORECASE,
                    )

                    degree = degree_part.strip()
                    field_of_study = field_part.strip()
                else:
                    degree = degree_line

                continue

        # First non-degree/non-date/non-grade line = institution
        if institution is None:
            institution = line
            continue

        # Possible location
        if location is None:
            location = line
            continue

    flush_education()

    return education


def _parse_experience(section_text: str) -> List[ResumeExperience]:
    """
    Conservative parser for work experience.

    Each detected block generally represents one position.
    """
    if not section_text.strip():
        return []

    lines = [
        _clean_line(line)
        for line in section_text.splitlines()
        if _clean_line(line)
    ]

    experiences = []

    current_company = None
    current_role = None
    current_start = None
    current_end = None
    current_bullets = []

    date_pattern = re.compile(
        r"([A-Za-z]+\s+\d{4})\s*[–-]\s*([A-Za-z]+\s+\d{4}|Present|Current)",
        re.IGNORECASE,
    )

    def flush_experience():
        nonlocal current_company
        nonlocal current_role
        nonlocal current_start
        nonlocal current_end
        nonlocal current_bullets

        if not current_company and not current_role:
            return

        experiences.append(
            ResumeExperience(
                company=current_company,
                role=current_role,
                start_date=current_start,
                end_date=current_end,
                bullets=current_bullets.copy(),
            )
        )

        current_company = None
        current_role = None
        current_start = None
        current_end = None
        current_bullets = []

    for line in lines:
        date_match = date_pattern.search(line)

        if date_match:
            if current_company and current_bullets:
                flush_experience()

            current_start = date_match.group(1)
            current_end = date_match.group(2)

            clean_without_date = date_pattern.sub("", line).strip()

            if clean_without_date:
                if current_role is None:
                    current_role = clean_without_date

            continue

        if re.match(r"^[•●▪◦‣*-]\s*", line):
            current_bullets.append(
                re.sub(r"^[•●▪◦‣*-]\s*", "", line)
            )
            continue

        # If we don't have a company yet, treat the line as company.
        if current_company is None:
            current_company = line
            continue

        # Next non-bullet line is normally the role.
        if current_role is None:
            current_role = line
            continue

    flush_experience()

    return experiences


def _parse_projects(section_text: str) -> List[ResumeProject]:
    """
    Initial project parser.

    Supports common:
        Project Name | Tech1, Tech2
    formats.
    """
    if not section_text.strip():
        return []

    lines = [
        _clean_line(line)
        for line in section_text.splitlines()
        if _clean_line(line)
    ]

    projects = []

    current_project = None
    current_technologies = []
    current_bullets = []

    def flush_project():
        nonlocal current_project
        nonlocal current_technologies
        nonlocal current_bullets

        if not current_project:
            return

        description = None

        if current_bullets:
            description = current_bullets[0]

        projects.append(
            ResumeProject(
                name=current_project,
                technologies=current_technologies.copy(),
                description=description,
                bullets=current_bullets.copy(),
                url=None,
            )
        )

        current_project = None
        current_technologies = []
        current_bullets = []

    for line in lines:
        if "|" in line:
            if current_project:
                flush_project()

            parts = [part.strip() for part in line.split("|")]

            current_project = parts[0]

            if len(parts) > 1:
                current_technologies = [
                    tech.strip()
                    for tech in re.split(r",|;", parts[1])
                    if tech.strip()
                ]

            continue

        if re.match(r"^[•●▪◦‣*-]\s*", line):
            current_bullets.append(
                re.sub(r"^[•●▪◦‣*-]\s*", "", line)
            )
            continue

        if current_project is None:
            current_project = line
        else:
            current_bullets.append(line)

    flush_project()

    return projects


def _parse_certifications(section_text: str) -> List[ResumeCertification]:
    if not section_text.strip():
        return []

    certifications = []

    for line in section_text.splitlines():
        line = _clean_line(line)

        if not line:
            continue

        line = re.sub(r"^[•●▪◦‣*-]\s*", "", line)

        issuer = None

        if " - " in line:
            name, issuer = line.split(" - ", 1)
        elif " | " in line:
            name, issuer = line.split(" | ", 1)
        else:
            name = line

        date_match = re.search(
            r"\b(19|20)\d{2}\b",
            line,
        )

        date = date_match.group(0) if date_match else None

        certifications.append(
            ResumeCertification(
                name=name.strip(),
                issuer=issuer.strip() if issuer else None,
                date=date,
                url=None,
            )
        )

    return certifications


def _parse_achievements(section_text: str) -> List[ResumeAchievement]:
    if not section_text.strip():
        return []

    achievements = []

    for line in section_text.splitlines():
        line = _clean_line(line)

        if not line:
            continue

        line = re.sub(r"^[•●▪◦‣*-]\s*", "", line)

        achievements.append(
            ResumeAchievement(
                title=line,
                description=None,
                date=None,
            )
        )

    return achievements


def parse_resume_profile(text: str) -> ResumeProfile:
    """
    Convert extracted resume text into a structured ResumeProfile.

    This function does not perform ATS scoring.
    It only creates a normalized representation of the resume.
    """
    if not text or not text.strip():
        raise ValueError("Resume text cannot be empty.")

    sections = _split_resume_sections(text)

    contact = _extract_contact(text)

    summary = sections["summary"] or None

    profile = ResumeProfile(
        contact=contact,
        summary=summary,
        education=_parse_education(sections["education"]),
        experience=_parse_experience(sections["experience"]),
        projects=_parse_projects(sections["projects"]),
        skills=_parse_skills(sections["skills"]),
        certifications=_parse_certifications(
            sections["certifications"]
        ),
        achievements=_parse_achievements(
            sections["achievements"]
        ),
    )

    return profile